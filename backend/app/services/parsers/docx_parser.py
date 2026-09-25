import io
import re
from typing import List
import docx
from app.services.parsers.base import BaseParser, ParsedContent


class DocxParser(BaseParser):
    def parse(self, file_bytes: bytes, filename: str = "") -> List[ParsedContent]:
        results: List[ParsedContent] = []
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        current_section = "Introduction"
        current_paragraphs: List[str] = []
        section_idx = 1

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""
            if "heading" in style_name.lower():
                # Flush previous section paragraphs if any
                if current_paragraphs:
                    section_text = "\n\n".join(current_paragraphs)
                    results.append(
                        ParsedContent(
                            text=section_text,
                            page_number=section_idx,
                            section=current_section,
                            metadata={"source": filename, "file_type": "docx"},
                        )
                    )
                    current_paragraphs = []
                    section_idx += 1
                current_section = text
            else:
                current_paragraphs.append(text)

        # Flush remaining section
        if current_paragraphs:
            section_text = "\n\n".join(current_paragraphs)
            results.append(
                ParsedContent(
                    text=section_text,
                    page_number=section_idx,
                    section=current_section,
                    metadata={"source": filename, "file_type": "docx"},
                )
            )

        # Fallback if no sections were parsed
        if not results and doc.paragraphs:
            all_text = "\n\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
            if all_text:
                results.append(
                    ParsedContent(
                        text=all_text,
                        page_number=1,
                        section="Document Body",
                        metadata={"source": filename, "file_type": "docx"},
                    )
                )

        return results
